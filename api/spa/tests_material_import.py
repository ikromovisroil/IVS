"""Materiallarni Excel'dan import qilish (api/spa/material_import.py).

    venv/Scripts/python.exe manage.py test api.spa.tests_material_import
"""
import io
import json

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from main.models import Material, MaterialCategory, MaterialMovement, Unit
from main.testing import make_employee, make_org, make_region

URL = "/api/v1/spa/warehouse/materials/import/"


def book(*sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile("import.xlsx", out.getvalue(),
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


TEMPLATE = ["Nomi", "O'lchov birligi", "1C kodi", "Narxi", "Soni"]
# inventarizatsiya ro'yxati (ИНВЕНТАРИЗАЦИОННАЯ ОПИСЬ) shakli: ikki qatorli sarlavha, raqamlash qatori, jami qatori
INVENTORY = [
    ["Канцтовары"],
    ["№", "Товарно-материал. Ценности", "Единица измерения", None, "Цена \nсум", "Фактическое наличие", None, "По данным бухгалтерского\nучета", None],
    [None, "наименование,\nвид, сорт, группа", "наименование", "код по\nСОЕИ", None, "кол-во", "сумма, сум", "кол-во", "сумма, сум"],
    [1, 2, 4, 5, 6, 7, 8, 9, 10],
    [1, "Блокнот #44252110#", "Штука", None, 25000, 100, 2500000, 100, 2500000],
    [2, "Ручка", "ед.", None, 5000, 10, 50000, 10, 50000],
    [3, "Пленка", "Пачка", None, 54870, 0.75, 41152.5, 0.75, 41152.5],
    [None, None, None, None, "Итого", 110.75, 2591152.5, 110.75, 2591152.5],
]


class MaterialImportTests(TestCase):
    def setUp(self):
        self.org = make_org("worker")
        self.region = make_region()
        self.root_user, self.root = make_employee(self.org, self.region, superuser=True)
        _, self.owner = make_employee(self.org, self.region, last="Javobgar")
        # bazadagi o'lchov birliklari (lotincha)
        self.dona, self.pachka, self.oram, self.kv = (
            Unit.objects.create(name=n) for n in ("Dona", "Pachka", "O'ram", "kv. metr"))
        self.api = APIClient()
        self.api.force_authenticate(self.root_user)

    def post(self, upload, **extra):
        data = {"file": upload, "organization": self.org.id, "employee": self.owner.id, **extra}
        return self.api.post(URL, data, format="multipart")

    def test_superuser_only(self):
        user, _ = make_employee(self.org, self.region)
        api = APIClient()
        api.force_authenticate(user)
        self.assertEqual(api.post(URL, {}, format="multipart").status_code, 403)
        self.assertEqual(api.get(URL + "options/").status_code, 403)
        self.assertEqual(api.get(URL + "template/").status_code, 403)

    def test_dry_run_writes_nothing(self):
        res = self.post(book(("Materiallar", [TEMPLATE, ["Ruchka", "Штука", "", 5000, 10]])))
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.data["dry_run"])
        self.assertEqual((res.data["created"], res.data["quantity"]), (1, 10))
        self.assertFalse(Material.objects.exists())
        self.assertFalse(MaterialMovement.objects.exists())
        self.assertEqual(Unit.objects.count(), 4)        # yangi birlik yaratilmaydi

    def test_import_creates_material_and_income_movement(self):
        res = self.post(book(("Materiallar", [TEMPLATE, ["Ruchka", "Штука", "", 5000, 10], ["Qalam", "dona", "Q-1", 1200.5, 3]])),
                        dry_run="0")
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual((res.data["created"], res.data["skipped"]), (2, 0))
        ruchka = Material.objects.get(name="Ruchka")
        self.assertEqual((ruchka.number, str(ruchka.price), ruchka.employee_id, ruchka.organization_id, ruchka.unit.name),
                         (10, "5000.00", self.owner.id, self.org.id, "Dona"))     # "Штука" bazadagi "Dona" ga mos
        self.assertIsNone(ruchka.category)
        mv = MaterialMovement.objects.get(material=ruchka)
        self.assertEqual((mv.income, mv.status, mv.employee_id, mv.user_id), (10, "created", self.owner.id, self.root.id))
        self.assertEqual(Material.objects.get(code="Q-1").number, 3)

    def test_identical_rows_are_all_imported_separately(self):
        """Jadval to'liq import qilinadi: nomi, narxi, soni va kodi bir xil qatorlar ham alohida material bo'ladi."""
        row = ["Ruchka", "Штука", "K-1", 5000, 10]
        res = self.post(book(("M", [TEMPLATE, row, row, ["Ruchka", "Штука", "K-1", 5000, 10]])), dry_run="0")
        self.assertEqual((res.data["rows"], res.data["created"], res.data["skipped"]), (3, 3, 0))
        self.assertEqual(Material.objects.filter(name="Ruchka").count(), 3)
        self.assertEqual(MaterialMovement.objects.filter(status="created", income=10).count(), 3)
        # qayta yuklash ham yangi materiallar yaratadi (mavjudga qo'shilmaydi)
        res = self.post(book(("M", [TEMPLATE, row])), dry_run="0")
        self.assertEqual(Material.objects.filter(name="Ruchka").count(), 4)
        self.assertEqual(sum(Material.objects.filter(name="Ruchka").values_list("number", flat=True)), 40)

    def test_inventory_format_sheet_category_units_and_skips(self):
        res = self.post(book(("Канцтовары", INVENTORY)), dry_run="0")
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual((res.data["rows"], res.data["created"], res.data["skipped"]), (3, 2, 1))
        self.assertIn("Kasr miqdor", res.data["issues"][0]["message"])
        # varaq nomi kategoriya hisoblanmaydi va kategoriya yaratilmaydi
        self.assertEqual(res.data["unmatched_categories"], [])
        self.assertFalse(MaterialCategory.objects.exists())
        notebook = Material.objects.get(name="Блокнот #44252110#")
        self.assertEqual((notebook.number, str(notebook.price), notebook.category, notebook.unit.name),
                         (100, "25000.00", None, "Dona"))
        # "ед." ham "Штука" ham bazadagi "Dona" ga mos
        self.assertEqual(Material.objects.get(name="Ручка").unit_id, notebook.unit_id)
        self.assertEqual(Unit.objects.count(), 4)
        self.assertEqual(res.data["unmatched_units"], [])      # "Пачка" qatori kasr miqdor sababli o'tkazildi

    def test_unmatched_unit_is_asked_and_mapped_not_created(self):
        # ruscha nomlar bazadagi lotincha birliklarga o'zi mos keladi; "Банка" bazada yo'q — so'raladi
        rows = book(("M", [TEMPLATE, ["Banka", "Банка", "", 100, 2], ["Papka", "Пачка", "", 100, 2],
                           ["Qop", "Упаковка", "", 10, 1], ["Taxta", "Метр Квадрат", "", 10, 1]]))
        res = self.post(rows, dry_run="0", unit_map="{}")
        self.assertEqual(res.data["unmatched_units"], [{"name": "Банка", "count": 1}])
        self.assertEqual(Material.objects.get(name="Papka").unit_id, self.pachka.id)
        self.assertEqual(Material.objects.get(name="Qop").unit_id, self.oram.id)
        self.assertEqual(Material.objects.get(name="Taxta").unit_id, self.kv.id)
        # moslama berilmasa — birliksiz yaratiladi, yangi birlik yaratilmaydi
        self.assertIsNone(Material.objects.get(name="Banka").unit)
        self.assertEqual(Unit.objects.count(), 4)
        # moslama berilsa — tanlangan birlik
        box = Unit.objects.create(name="Quti")
        res = self.post(book(("M", [TEMPLATE, ["Banka2", "Банка", "", 100, 2]])), dry_run="0",
                        unit_map=json.dumps({"Банка": box.id}))
        self.assertEqual(res.data["unmatched_units"], [])
        self.assertEqual(Material.objects.get(name="Banka2").unit_id, box.id)
        bad = self.post(book(("M", [TEMPLATE, ["X", "", "", 1, 1]])), unit_map="[1")
        self.assertEqual(bad.status_code, 400)

    def test_category_only_existing_from_column_or_chosen(self):
        col_cat = MaterialCategory.objects.create(name="kantselyariya", organization=self.org)
        chosen = MaterialCategory.objects.create(name="Boshqa", organization=None)
        head = TEMPLATE + ["Kategoriya"]
        res = self.post(book(("Varaq", [head, ["Qalam", "dona", "", 1000, 2, "Kantselyariya"],
                                        ["Daftar", "dona", "", 1500, 1, "Yo'q kategoriya"], ["Klip", "dona", "", 10, 1, None]])), dry_run="0")
        # Excel ustunidagi nom bazadagi kategoriyaga mos keladi; topilmaganlar kategoriyasiz, yangisi yaratilmaydi
        self.assertEqual(Material.objects.get(name="Qalam").category_id, col_cat.id)
        self.assertIsNone(Material.objects.get(name="Daftar").category)
        self.assertEqual(res.data["unmatched_categories"], ["Yo'q kategoriya"])
        self.assertEqual(MaterialCategory.objects.count(), 2)
        # kategoriya aniq tanlansa — hamma qatorga shu
        res = self.post(book(("Xoz", [TEMPLATE, ["Qalam2", "dona", "", 1000, 2]])), dry_run="0", category=chosen.id)
        self.assertEqual(Material.objects.get(name="Qalam2").category_id, chosen.id)
        # boshqa tashkilotning kategoriyasi tanlanmaydi
        alien = MaterialCategory.objects.create(name="Alien", organization=make_org("worker"))
        res = self.post(book(("Xoz", [TEMPLATE, ["Qalam3", "dona", "", 1000, 2]])), category=alien.id)
        self.assertEqual(res.status_code, 400)

    def test_validation(self):
        bad = self.api.post(URL, {"file": SimpleUploadedFile("a.txt", b"x"), "organization": self.org.id, "employee": self.owner.id}, format="multipart")
        self.assertEqual(bad.status_code, 400)
        _, other = make_employee(make_org("worker"), self.region)
        res = self.post(book(("M", [TEMPLATE, ["A", "", "", 1, 1]])), employee=other.id)
        self.assertEqual(res.status_code, 400)
        self.assertIn("employee", res.data)
        res = self.post(book(("M", [["Hech nima"], ["A", 1]])))
        self.assertEqual(res.status_code, 400)
        self.assertIn("file", res.data)

    def test_options_and_template(self):
        res = self.api.get(URL + "options/", {"organization": self.org.id})
        self.assertEqual(res.status_code, 200)
        self.assertIn(self.owner.id, [e["id"] for e in res.data["employees"]])
        self.assertIn(self.org.id, [o["id"] for o in res.data["organizations"]])
        self.assertEqual(sorted(u["name"] for u in res.data["units"]), ["Dona", "O'ram", "Pachka", "kv. metr"])
        cat = MaterialCategory.objects.create(name="Mening", organization=self.org)
        MaterialCategory.objects.create(name="Begona", organization=make_org("worker"))
        res = self.api.get(URL + "options/", {"organization": self.org.id})
        self.assertEqual([c["id"] for c in res.data["categories"]], [cat.id])
        tpl = self.api.get(URL + "template/")
        self.assertEqual(tpl.status_code, 200)
        self.assertIn("spreadsheetml", tpl["Content-Type"])
        # namuna fayl o'zi import qilinadi
        upload = SimpleUploadedFile("t.xlsx", tpl.content)
        res = self.post(upload, dry_run="1")
        self.assertEqual((res.status_code, res.data["rows"]), (200, 1))
        # namuna: Nomi | O'lchov birligi | 1C kodi | Narxi | Soni, bitta misol qator
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(tpl.content)).active
        self.assertEqual([c.value for c in ws[1]], ["Nomi", "O'lchov birligi", "1C kodi", "Narxi", "Soni"])
        self.assertEqual([c.value for c in ws[2]], ["Ruchka kanselyariya", "dona", None, 5000, 100])
        self.assertEqual(res.data["unmatched_units"], [])        # "dona" bazadagi "Dona" ga mos
