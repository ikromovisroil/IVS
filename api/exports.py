"""
Excel (.xlsx) eksportlar — mobil ilova uchun (JWT).

* employees.xlsx va material-report.xlsx — saytdagi export_employee_xlsx / export_mat_info_xlsx ni to'g'ridan-to'g'ri
  chaqiradi (ularda ruxsat va ko'rish doirasi allaqachon bor).
* technics.xlsx va materials.xlsx — saytdagi eksportlar tashkilot/hudud bo'yicha CHEKLANMAGAN (kirgan har kim hammasini
  yuklay oladi); bu yerda API ro'yxatlari bilan bir xil huquq va ko'rish doirasi qo'llanadi.
"""
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from rest_framework import permissions
from rest_framework.exceptions import PermissionDenied
from rest_framework.views import APIView

from main.ajax_xlsx import export_employee_xlsx, export_mat_info_xlsx
from main.models import Technics

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_response(wb, filename, width):
    ws = wb.active
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    bio = BytesIO()
    wb.save(bio)
    resp = HttpResponse(bio.getvalue(), content_type=XLSX_TYPE)
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


class _ExportBase(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def employee(self, request):
        emp = getattr(request.user, "employee", None)
        if not emp:
            raise PermissionDenied("Employee yo'q")
        return emp


class _BridgeExportView(_ExportBase):
    """Saytdagi (Django) eksport funksiyasini JWT foydalanuvchi bilan chaqiradi."""
    web_view = None

    def get(self, request):
        django_request = request._request
        django_request.user = request.user
        return self.web_view(django_request)


class EmployeesExportView(_BridgeExportView):
    """GET /api/export/employees.xlsx?organization=&region=&department=&directorate=&division=&name= (view_employee; organization majburiy)."""
    web_view = staticmethod(export_employee_xlsx)


class MaterialReportExportView(_BridgeExportView):
    """GET /api/export/material-report.xlsx?employee=&date1=&date2=&name= — kirim-chiqim hisoboti (view_material)."""
    web_view = staticmethod(export_mat_info_xlsx)


class TechnicsExportView(_ExportBase):
    """GET /api/export/technics.xlsx — texnikalar (view_technics; ro'yxatdagi kabi tashkilot/hudud doirasi). Filtrlar ro'yxatniki bilan bir xil."""

    def get(self, request):
        from .views import TechnicsViewSet

        if not request.user.has_perm("main.view_technics"):
            raise PermissionDenied("Sizga ruxsat yo'q")

        view = TechnicsViewSet(request=request, action="list", format_kwarg=None)
        qs = view.filter_queryset(view.get_queryset()).select_related(
            "employee__organization", "employee__department", "employee__directorate", "employee__division",
        ).prefetch_related("structure_set")

        wb = Workbook()
        ws = wb.active
        ws.title = "Technics"
        ws.append([
            "Tashkilot", "Departament", "Boshqarma", "Bo'lim", "F.I.O",
            "Group", "Category", "Name", "Parametr", "I/N", "S/N", "Mac",
            "Status", "Manitor Name", "Manitor S/N",
        ])
        for t in qs.order_by("-id"):
            emp = t.employee
            monitor = next(iter(s for s in t.structure_set.all() if s.is_active), None)
            ws.append([
                emp.organization.name if emp and emp.organization_id else "",
                emp.department.name if emp and emp.department_id else "",
                emp.directorate.name if emp and emp.directorate_id else "",
                emp.division.name if emp and emp.division_id else "",
                emp.full_name if emp else "",
                t.group.name if t.group_id else "",
                t.category.name if t.category_id else "",
                t.name or "", t.parametr or "", t.inventory or "", t.serial or "", t.mac or "",
                t.get_status_display(),
                monitor.name if monitor else "", monitor.serial if monitor else "",
            ])
        return _xlsx_response(wb, "technics.xlsx", 20)


class MaterialsExportView(_ExportBase):
    """GET /api/export/materials.xlsx — materiallar (view_material; ro'yxatdagi kabi ko'rish doirasi). Filtrlar: employee, category, unit, search."""

    def get(self, request):
        from .views import MaterialViewSet

        if not request.user.has_perm("main.view_material"):
            raise PermissionDenied("Sizga ruxsat yo'q")

        view = MaterialViewSet(request=request, action="list", format_kwarg=None)
        qs = view.filter_queryset(view.get_queryset()).select_related("employee", "unit", "category")

        wb = Workbook()
        ws = wb.active
        ws.title = "Material"
        ws.append(["Xodim", "Material nomi", "Soni", "Narxi", "Summa", "O'lchov birligi", "1C kod"])
        for m in qs.order_by("-id"):
            price = m.price or 0
            ws.append([
                m.employee.full_name if m.employee_id else "",
                m.name or "", m.number, price, (m.number or 0) * price,
                m.unit.name if m.unit_id else "", m.code or "",
            ])
        return _xlsx_response(wb, "material.xlsx", 22)
